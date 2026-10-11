"""Independent validation of saved outputs. Usage: validate.py OUT_PROJECT_DIR FIXTURE_DIR"""
import json, os, sys, subprocess, zipfile, gzip, io, re, glob
from PIL import Image
root, FX = sys.argv[1], sys.argv[2]
res = []
def sh(*a, inp=None):
    r = subprocess.run(a, capture_output=True, input=inp)
    return r.returncode, r.stdout.decode('utf-8', 'replace'), r.stderr.decode('utf-8', 'replace')
def out(case, i=0):
    d = os.path.join(root, case)
    files = sorted(f for f in os.listdir(d) if f != 'record.json')
    return os.path.join(d, files[i]) if len(files) > i else None
def outs(case):
    d = os.path.join(root, case)
    return [os.path.join(d, f) for f in sorted(os.listdir(d)) if f != 'record.json']
def rec(case): return json.load(open(os.path.join(root, case, 'record.json')))
def check(case, ok, detail):
    res.append((case, 'PASS' if ok else 'FAIL', detail))
def near(a, b, tol=12): return all(abs(x - y) <= tol for x, y in zip(a, b))

def quad_check(case, alpha_expected, tol=12, path=None):
    p = path or out(case)
    im = Image.open(p); im.load(); fmt = im.format
    im = im.convert('RGBA'); w, h = im.size
    px = lambda x, y: im.getpixel((int(x * w), int(y * h)))
    tl, tr, bl, br = px(.25, .25), px(.75, .25), px(.25, .75), px(.75, .75)
    ok = near(tl[:3], (255, 0, 0), tol) and near(tr[:3], (0, 255, 0), tol) and near(bl[:3], (0, 0, 255), tol)
    if alpha_expected: ok = ok and br[3] < 10 and tl[3] > 245
    else: ok = ok and br[3] == 255
    check(case, ok and (w, h) == (400, 300), f'{fmt} {w}x{h} TL{tl} TR{tr} BL{bl} BR{br}')

for t, a in [('webp', True), ('png', True), ('jpeg', False), ('avif', True), ('gif', True), ('bmp', False), ('tiff', True), ('pdf', None), ('ico', True), ('svg-wrapper', None)]:
    case = f'quad.png__image-to-{t}'
    if not os.path.isdir(os.path.join(root, case)): continue
    p = out(case)
    if t == 'avif':
        # Pillow 12 decodes AVIF when built with libavif; fall back to ffmpeg.
        try: quad_check(case, True, tol=20)
        except Exception as e:
            rc, o, er = sh('ffmpeg', '-v', 'error', '-i', p, '-f', 'null', '-'); check(case, rc == 0, f'pillow failed ({e}); ffmpeg decode rc={rc} {er[:200]}')
    elif t == 'pdf':
        rc, o, _ = sh('pdfinfo', p); rc2, _, e2 = sh('qpdf', '--check', p)
        # Rendered at 96 DPI so the 300x225 pt page becomes 400x300 px; transparency is flattened to white in a PDF.
        rc3, _, _ = sh('pdftoppm', '-r', '96', '-png', '-singlefile', p, p + '.render')
        quad_check(case + ' (rendered)', False, tol=20, path=p + '.render.png') if rc3 == 0 else None
        check(case, rc == 0 and rc2 == 0 and re.search(r'Pages:\s+1\n', o) is not None, re.sub(r'\s+', ' ', o)[:200] + ' qpdf=' + str(rc2))
    elif t == 'ico':
        im = Image.open(p); check(case, im.format == 'ICO', f'ICO sizes={im.info.get("sizes")} size={im.size}')
        im.size = max(im.info.get('sizes', [im.size])); 
        tmp = p + '.png'; im.convert('RGBA').resize((400, 300)).save(tmp); quad_check(case + ' (largest icon, resized)', True, tol=25, path=tmp)
    elif t == 'svg-wrapper':
        s = open(p).read(); m = re.search(r'data:image/(\w+);base64,([A-Za-z0-9+/=]+)', s)
        import base64
        if m:
            tmp = p + '.inner'; open(tmp, 'wb').write(base64.b64decode(m.group(2))); quad_check(case + ' (embedded bitmap)', True, path=tmp)
        check(case, s.lstrip().startswith(('<svg', '<?xml')) and bool(m), f'svg root ok, embedded={m.group(1) if m else None}')
    else:
        quad_check(case, a, tol=20 if t in ('gif',) else 12)

for src in ['bmp', 'tiff', 'webp', 'ico']:
    case = f'quad.{src}__image-to-png'
    if src == 'ico':
        # The source ICO is 64x48 (Pillow keeps the 4:3 aspect when writing a 64 px icon).
        im = Image.open(out(case)); check(case, im.size == (64, 48), f'{im.size}')
    else: quad_check(case, src != 'bmp')
quad_check('really-png.pdf__image-to-jpeg', False)

# Animated GIF -> first frame red
for case in ['anim.gif__image-to-png', 'anim.gif__image-to-webp']:
    im = Image.open(out(case)).convert('RGB'); c = im.getpixel((60, 40)); check(case, near(c, (255, 0, 0)) and im.size == (120, 80), f'{im.size} centre={c}')

# Orientation 6: stored 400x200 (left red / right blue); displayed should be 200x400 with top red, bottom blue
for case in ['oriented-gps.jpg__image-to-png', 'oriented-gps.jpg__image-to-webp']:
    im = Image.open(out(case)).convert('RGB'); w, h = im.size
    top, bot = im.getpixel((w // 2, h // 4)), im.getpixel((w // 2, 3 * h // 4))
    rc, o, _ = sh('exiftool', '-j', '-a', '-G1', out(case)); tags = json.loads(o)[0]
    leaked = [k for k in tags if any(s in k for s in ('GPS', 'Artist', 'Make'))]
    check(case, (w, h) == (200, 400) and near(top, (255, 0, 0), 20) and near(bot, (0, 0, 255), 20) and not leaked, f'{w}x{h} top={top} bottom={bot} leaked={leaked}')
case = 'oriented-gps.jpg__image-strip-metadata'
rc, o, _ = sh('exiftool', '-j', '-a', '-G1', out(case)); tags = json.loads(o)[0]
leaked = {k: v for k, v in tags.items() if any(s in k for s in ('GPS', 'Artist', 'Make'))}
check(case, not leaked and tags.get('IFD0:Orientation') in ('Rotate 90 CW', None), f'leaked={leaked} orientation={tags.get("IFD0:Orientation")}')
src_px = Image.open(os.path.join(FX, 'oriented-gps.jpg')).tobytes(); out_px = Image.open(out(case)).tobytes()
check(case + ' (pixels identical)', src_px == out_px, f'lossless={src_px == out_px}')
case = 'oriented-gps.jpg__image-to-pdf'
rc, o, _ = sh('pdfinfo', out(case)); check(case, rc == 0 and re.search(r'Pages:\s+1\n', o) is not None, re.search(r'Page size:.*', o).group(0) if rc == 0 else o)
sh('pdftoppm', '-r', '72', '-png', '-singlefile', out(case), out(case) + '.render'); im = Image.open(out(case) + '.render.png').convert('RGB'); w, h = im.size
check(case + ' (orientation)', h > w and near(im.getpixel((w // 2, h // 4)), (255, 0, 0), 30), f'rendered {w}x{h} top={im.getpixel((w // 2, h // 4))}')

im = Image.open(out('remote.svg__image-to-png')).convert('RGB'); check('remote.svg__image-to-png', near(im.getpixel((50, 25)), (0, 255, 0)) and im.size == (100, 50), f'{im.size} {im.getpixel((50, 25))}')
im = Image.open(out('photo.heic__image-to-jpeg')); check('photo.heic__image-to-jpeg', im.format == 'JPEG', f'{im.format} {im.size}')
t = open(out('sign.png__image-ocr'), encoding='utf-8').read(); check('sign.png__image-ocr', len(t.strip()) > 3, repr(t[:80]))
t = open(out('scanned.pdf__pdf-ocr'), encoding='utf-8').read(); check('scanned.pdf__pdf-ocr', 'total amount due' in t.lower(), repr(t[:90]))

# Documents
def pdftext(p): return sh('pdftotext', '-layout', p, '-')[1]
def pdfok(case, needles):
    for p in outs(case):
        rc, o, _ = sh('qpdf', '--check', p); t = pdftext(p)
        miss = [n for n in needles if n not in t]
        check(case, rc == 0 and not miss, f'qpdf={rc} missing={miss} text={t.strip()[:80]!r}')
h = open(out('doc.md__md-to-html'), encoding='utf-8').read()
bad = [s for s in ['<script', 'tracker.example', 'javascript:'] if s in h]
check('doc.md__md-to-html', '<h1' in h and '<table' in h and not bad, f'remote/script refs left: {bad}; ' + '; '.join(re.findall(r'<[^>]*tracker\.example[^>]*>', h))[:300])
t = open(out('doc.md__md-to-txt'), encoding='utf-8').read(); check('doc.md__md-to-txt', 'Audit Heading' in t and 'Ελληνικά' in t and '<script' not in t, repr(t[:120]))
m = open(out('page.html__html-to-md'), encoding='utf-8').read(); check('page.html__html-to-md', '# HTML Heading' in m and '| H1 | H2 |' in m and 'tracker' not in m and 'javascript:' not in m, repr(m[:300]))
pdfok('page.html__html-to-pdf', ['HTML Heading', 'item one', 'c1'])
t = open(out('page.html__html-to-txt'), encoding='utf-8').read(); check('page.html__html-to-txt', 'HTML Heading' in t and 'alert' not in t, repr(t[:120]))
pdfok('plain.txt__txt-to-pdf', ['Audit plain text line one', 'Line two'])
h = open(out('plain.txt__txt-to-html'), encoding='utf-8').read(); check('plain.txt__txt-to-html', 'Audit plain text line one' in h, h[-200:].replace('\n', ' '))
h = open(out('html-named.txt__txt-to-html'), encoding='utf-8').read(); check('html-named.txt__txt-to-html', '<script>' not in h and '&lt;script&gt;' in h, 'escaped' if '&lt;script' in h else 'NOT ESCAPED')
b = open(out('cp1252.txt__txt-reencode'), 'rb').read(); check('cp1252.txt__txt-reencode', b.decode('utf-8').startswith('Café naïve – “quotes” €100'), repr(b[:60]))
g = gzip.decompress(open(out('plain.txt__file-to-gzip'), 'rb').read()); check('plain.txt__file-to-gzip', g == open(os.path.join(FX, 'plain.txt'), 'rb').read(), 'roundtrip')
t = open(out('three.pdf__pdf-to-txt'), encoding='utf-8').read(); check('three.pdf__pdf-to-txt', all(f'PAGE{i}MARK' in t for i in (1, 2, 3)), repr(t[:140]))
for kind in ['png', 'jpeg']:
    case = f'three.pdf__pdf-to-{kind}'; ps = outs(case)
    check(case, len(ps) == 3 and all(Image.open(p).size[0] > 100 for p in ps), f'{len(ps)} pages, sizes={[Image.open(p).size for p in ps]}')
p = out('three.pdf__pdf-extract-pages'); rc, o, _ = sh('pdfinfo', p); check('three.pdf__pdf-extract-pages', re.search(r'Pages:\s+3\n', o) is not None and not os.path.basename(p).endswith('-.pdf'), f'name={os.path.basename(p)} ' + re.search(r'Pages:\s+\d+', o).group(0))
p = out('three.pdf__pdf-strip-metadata'); rc, o, _ = sh('exiftool', '-j', '-a', p); tags = json.loads(o)[0]
raw = open(p, 'rb').read()
check('three.pdf__pdf-strip-metadata', 'Secret Author' not in json.dumps(tags) and b'Secret Author' not in raw and all(f'PAGE{i}MARK' in pdftext(p) for i in (1, 2, 3)), f'Author tag={tags.get("Author")} Producer={tags.get("Producer")} Creator={tags.get("Creator")} raw-bytes-has-author={b"Secret Author" in raw}')
h = open(out('doc.docx__docx-to-html'), encoding='utf-8').read(); check('doc.docx__docx-to-html', 'Docx Heading' in h and '<table' in h and 'data:image' in h, 'heading/table/img present' if 'data:image' in h else 'missing image')
m = open(out('doc.docx__docx-to-md'), encoding='utf-8').read(); check('doc.docx__docx-to-md', 'Docx Heading' in m and 'T1' in m, repr(m[:200]))
t = open(out('doc.docx__docx-to-txt'), encoding='utf-8').read(); check('doc.docx__docx-to-txt', 'Docx paragraph with Ünïcode.' in t, repr(t[:100]))
pdfok('doc.docx__docx-to-pdf', ['Docx Heading', 'Ünïcode', 'T2'])
im = Image.open(out('doc.docx__docx-extract-media')).convert('RGB'); check('doc.docx__docx-extract-media', im.size == (40, 30) and near(im.getpixel((5, 5)), (255, 0, 0)), f'{im.size}')
t = open(out('deck.pptx__pptx-to-txt'), encoding='utf-8').read(); check('deck.pptx__pptx-to-txt', t.find('Slide 1 title') < t.find('Slide 2 title') and 'Notes 2' in t and t.find('Slide 1') >= 0, repr(t[:200]))
im = Image.open(out('deck.pptx__pptx-extract-media')); check('deck.pptx__pptx-extract-media', im.size == (40, 30), f'{im.size}')

# Spreadsheets
import openpyxl, csv as csvm
c = open(out('inject.csv__sheet-to-csv'), encoding='utf-8-sig').read(); rows = list(csvm.reader(io.StringIO(c)))
dang = [v for r in rows for v in r if v[:1] in '=+-@']
check('inject.csv__sheet-to-csv', rows[1][0] == '02134' and rows[2][1] == 'multi\nline' and not dang, f'rows={rows} dangerous={dang}')
wb = openpyxl.load_workbook(out('inject.csv__sheet-to-xlsx')); ws = wb.active
vals = [[c.value for c in r] for r in ws.iter_rows()]; types = [[c.data_type for c in r] for r in ws.iter_rows()]
check('inject.csv__sheet-to-xlsx', vals[1][0] == '02134' and not any(t == 'f' for r in types for t in r), f'values={vals} types={types}')
j = json.load(open(out('inject.csv__sheet-to-json'))); check('inject.csv__sheet-to-json', j[0]['id'] == '02134' and j[0]['formula'] == '=1+2', json.dumps(j)[:200])
t = open(out('inject.csv__sheet-to-tsv'), encoding='utf-8-sig').read(); check('inject.csv__sheet-to-tsv', '02134' in t and '\t' in t, repr(t[:120]))
z = zipfile.ZipFile(out('inject.csv__sheet-to-ods')); check('inject.csv__sheet-to-ods', z.read('mimetype') == b'application/vnd.oasis.opendocument.spreadsheet' and b'02134' in z.read('content.xml'), z.namelist()[0])
c = open(out('book.xlsx__sheet-to-csv'), encoding='utf-8-sig').read(); check('book.xlsx__sheet-to-csv', '02134' in c and 'Widget' in c, repr(c))
j = json.load(open(out('book.xlsx__sheet-to-json'))); check('book.xlsx__sheet-to-json', 'Widget' in json.dumps(j), json.dumps(j)[:300])
r = rec('book.xlsx__sheet-to-csv'); check('book.xlsx__sheet-to-csv (uncomputed-formula notice)', 'formula' in r['cardText'].lower(), r['cardText'].replace('\n', ' | ')[:300])
# Formulas with no cached result cannot be written to ODS; the app must say so rather than drop them silently.
z = zipfile.ZipFile(out('book.xlsx__sheet-to-ods')); cx = z.read('content.xml').decode(); warned = 'left empty' in rec('book.xlsx__sheet-to-ods')['cardText']
check('book.xlsx__sheet-to-ods', 'Sales' in cx and 'Notes' in cx and ('of:=' in cx or warned), f'both sheets; formula kept={"of:=" in cx}; uncached-formula warning shown={warned}')
if os.path.isdir(os.path.join(root, 'book-cached.xlsx__sheet-to-ods')):
    cx = zipfile.ZipFile(out('book-cached.xlsx__sheet-to-ods')).read('content.xml').decode()
    check('book-cached.xlsx__sheet-to-ods', 'of:=' in cx and '7.5' in cx, 'formula kept' if 'of:=' in cx else 'formula missing')
    wb = openpyxl.load_workbook(out('book-cached.xlsx__sheet-to-xlsx')); f = wb['Sales']['E2'].value
    check('book-cached.xlsx__sheet-to-xlsx', isinstance(f, str) and f.startswith('=') and len(wb.sheetnames) == 2, f'E2={f!r} sheets={wb.sheetnames}')
    c = open(out('book-cached.xlsx__sheet-to-csv'), encoding='utf-8-sig').read(); check('book-cached.xlsx__sheet-to-csv', '7.5' in c and '20' in c, repr(c))
c = open(out('rows.json__sheet-to-csv'), encoding='utf-8-sig').read(); check('rows.json__sheet-to-csv', '02134' in c and "'=HYPERLINK" in c, repr(c))
wb = openpyxl.load_workbook(out('rows.json__sheet-to-xlsx')); vals = [[c.value for c in r] for r in wb.active.iter_rows()]; types = [[c.data_type for c in r] for r in wb.active.iter_rows()]
check('rows.json__sheet-to-xlsx', not any(t == 'f' for r in types for t in r), f'{vals} {types}')
import yaml as _y  # noqa
y = _y.safe_load(open(out('data.json__data-to-yaml'))); check('data.json__data-to-yaml', y == json.load(open(os.path.join(FX, 'data.json'))), 'roundtrip equal' if y == json.load(open(os.path.join(FX, 'data.json'))) else f'DIFF {y}')
x = open(out('data.json__data-to-xml')).read(); rc, o, e = sh('python3', '-I', '-c', 'import sys,xml.dom.minidom as m;m.parse(sys.argv[1])', out('data.json__data-to-xml'))
check('data.json__data-to-xml', rc == 0 and 'Audit &amp; Co' in x and '02134' in x, f'well-formed={rc == 0} {x[:200]!r}')
j = json.load(open(out('conf.yaml__data-to-json'))); check('conf.yaml__data-to-json', j == _y.safe_load(open(os.path.join(FX, 'conf.yaml'))), json.dumps(j))
j = json.load(open(out('data.xml__data-to-json'))); s = json.dumps(j)
check('data.xml__data-to-json', 'XML & You' in s, s[:300])
check('data.xml__data-to-json (leading-zero value kept)', '"02134"' in s, 'zip values: ' + str(re.findall(r'"zip": [^,}]*', s)))
v = open(out('subs.srt__subs-to-vtt'), encoding='utf-8').read(); check('subs.srt__subs-to-vtt', v.startswith('WEBVTT') and '00:00:01.000 --> 00:00:02.500' in v and 'Ünïcode 日本' in v, repr(v))

# Archives
def same(p, s): return open(p, 'rb').read() == s
ps = outs('nested.zip__archive-extract'); names = [os.path.basename(p)[3:] for p in ps]
check('nested.zip__archive-extract', same(ps[0], b'alpha\n') and same(ps[1], b'bravo\n') and not any('/' in n or n.startswith('.') for n in names), f'{names}')
ps = outs('nested.tar.gz__archive-extract'); names = [os.path.basename(p)[3:] for p in ps]
check('nested.tar.gz__archive-extract', len(ps) == 2 and same(ps[0], b'alpha\n') and same(ps[1], b'bravo\n'), f'{names} (symlink skipped={"link" not in names})')
p = out('nested.tar.gz__archive-to-zip'); rc, o, _ = sh('unzip', '-l', p); rc2, _, _ = sh('unzip', '-tq', p)
check('nested.tar.gz__archive-to-zip', rc2 == 0 and 'sub/b.txt' in o and 'link' not in o.split(), re.sub(r'\s+', ' ', o)[:200])
ps = outs('plain.tar__archive-extract'); check('plain.tar__archive-extract', same(ps[0], b'alpha\n') and same(ps[1], b'bravo\n'), 'ok')
ps = outs('truncated.tar__archive-extract'); check('truncated.tar__archive-extract (must be refused: header declares 4096 B)', not ps, f'{len(ps)} output(s); size={os.path.getsize(ps[0]) if ps else 0}')
p = out('hello.txt.gz__archive-extract'); check('hello.txt.gz__archive-extract', same(p, b'gzip payload\n'), 'ok')
ps = outs('nested.7z__archive-extract-7z'); check('nested.7z__archive-extract-7z', same(ps[0], b'alpha\n') and same(ps[1], b'bravo\n'), 'ok')

# Media: ffprobe + full decode
def probe(p):
    rc, o, e = sh('ffprobe', '-v', 'error', '-show_streams', '-show_format', '-of', 'json', p)
    return json.loads(o) if rc == 0 else None
def full_decode(p):
    rc, o, e = sh('ffmpeg', '-v', 'error', '-i', p, '-f', 'null', '-'); return rc == 0 and not e.strip(), e.strip()[:150]
def media(case, want_codecs, dur=None, tol=0.25, ch=None):
    p = out(case); j = probe(p)
    if not j: return check(case, False, 'ffprobe failed')
    codecs = sorted(s['codec_name'] for s in j['streams']); d = float(j['format'].get('duration', 0) or 0)
    ok_dec, err = full_decode(p)
    a = [s for s in j['streams'] if s['codec_type'] == 'audio']
    ok = all(c in codecs for c in want_codecs) and ok_dec and (dur is None or abs(d - dur) <= tol) and (ch is None or (a and a[0].get('channels') == ch))
    tags = j['format'].get('tags', {})
    check(case, ok, f'{j["format"]["format_name"]} codecs={codecs} dur={d:.3f} ch={a[0].get("channels") if a else None} decode_clean={ok_dec} {err} tags={ {k: v for k, v in tags.items() if k.lower() in ("artist",)} }')
for t, c in [('mp3', ['mp3']), ('m4a', ['aac']), ('aac', ['aac']), ('ogg', ['vorbis']), ('opus', ['opus']), ('flac', ['flac']), ('aiff', ['pcm_s16be'])]:
    # ffprobe only estimates the duration of raw ADTS AAC from bitrate, so it gets a wider tolerance.
    media(f'stereo.wav__audio-to-{t}', c, 3.0, 0.15 if t == 'aac' else 0.1, 2)
media('tone.mp3__audio-to-wav', ['pcm_s16le'], 3.0, 0.1, 2); media('tone.flac__audio-to-mp3', ['mp3'], 3.0, 0.1, 2)
media('clip.mp4__video-to-webm', ['vp8', 'vorbis'], 2.0); media('clip.mp4__video-to-mov', ['h264', 'aac'], 2.0); media('clip.mp4__video-to-mkv', ['h264', 'aac'], 2.0)
media('clip.mp4__video-to-gif', ['gif'], 2.0, 0.3); media('clip.mp4__audio-to-m4a', ['aac'], 2.0); media('clip.webm__video-to-mp4', ['h264', 'aac'], 2.0)
im = Image.open(out('clip.mp4__video-thumbnail')); check('clip.mp4__video-thumbnail', im.format == 'JPEG' and im.size == (320, 240), f'{im.size}')

# JPEG streams embedded in PDFs must not carry GPS (the UI promises "Converting removes it").
def jpeg_streams(pdf):
    b = open(pdf, 'rb').read(); i = 0; found = []
    while (st := b.find(b'\xff\xd8\xff', i)) >= 0:
        e = b.find(b'\xff\xd9', st) + 2; found.append(b[st:e]); i = e
    return found
def gps_in_pdf(pdf):
    hits = []
    for n, j in enumerate(jpeg_streams(pdf)):
        tmp = f'{pdf}.stream{n}.jpg'; open(tmp, 'wb').write(j)
        rc, o, _ = sh('exiftool', '-j', '-a', '-G1', tmp); t = json.loads(o)[0]
        hits += [f'stream{n}:{k}={v}' for k, v in t.items() if k.startswith('GPS:') or k == 'IFD0:Make']
    return hits
for case in ['oriented-gps.jpg__image-to-pdf']:
    h = gps_in_pdf(out(case)); check(case + ' (embedded JPEG metadata)', not h, '; '.join(h)[:300] or 'clean')
rp = os.path.join(root, 'repro')
if os.path.isdir(rp):
    for name in ['combined.pdf', 'gps-upright.jpg.image-to-pdf.out']:
        f = os.path.join(rp, name)
        if os.path.exists(f): h = gps_in_pdf(f); check(f'repro {name} (embedded JPEG metadata)', not h, '; '.join(h)[:300] or 'clean')

for c, s, d in res: print(f'{s}  {c:60}  {d}')
print('TOTAL', len(res), 'PASS', sum(1 for r in res if r[1] == 'PASS'), 'FAIL', sum(1 for r in res if r[1] == 'FAIL'))
